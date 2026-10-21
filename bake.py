"""Official population figures for places the World Bank leaves out (Helen, 2026-10-10: "Taiwan has its own government and
population"; then option 4: each place's own official statistics, Natural Earth's estimate only where none exists).
One JSON per source (never merged: each keeps its own terms), plus index.json listing them:
  out/<id>.json  {"source", "publisher", "url", "licence", "places": [{"iso3", "name", "population", "aged_65_plus_pct",
                  "density_per_km2", "period"}]}      (a figure the source does not publish is left out, never estimated)
Usage: python3 bake.py [<source id> ...]     (default: every source)"""
import csv, datetime, html, io, json, os, re, subprocess, sys, time, urllib.request
import openpyxl
import sources

UA = {'User-Agent': 'Milleleaf world-places (https://github.com/milleleaf/world-places)'}


def get(url, data=None, headers=None, tries=1):
    req = urllib.request.Request(url, data=data, headers={**UA, **(headers or {})})
    for t in range(tries):
        try:
            with urllib.request.urlopen(req, timeout=120) as r:
                return r.read()
        except OSError:
            if t == tries - 1:
                raise
            time.sleep(10)


def xlsx(url, tries=1):
    return openpyxl.load_workbook(io.BytesIO(get(url, tries=tries)), read_only=True, data_only=True)


def text(page):
    """An HTML page as plain text (tags dropped, whitespace collapsed)."""
    page = re.sub(r'<(script|style)\b.*?</\1>', ' ', page, flags=re.S | re.I)
    return re.sub(r'\s+', ' ', html.unescape(re.sub(r'<[^>]+>', ' ', page))).strip()


def pct(old, total):
    return round(100 * old / total, 2)


def taiwan():
    """Ministry of the Interior, household registration: every village, by single year of age (ODRP014, monthly)."""
    now = datetime.date.today()
    for back in range(6):
        y, m = divmod(now.year * 12 + now.month - 1 - back, 12)
        ym = f'{y - 1911}{m + 1:02d}'
        url = lambda pg: f'https://www.ris.gov.tw/rs-opendata/api/v1/datastore/ODRP014/{ym}?page={pg}'
        first = json.loads(get(url(1)))
        if not first.get('responseData'):
            continue                                                       # that month is not out yet
        rows = first['responseData'] + [r for pg in range(2, int(first['totalPage']) + 1) for r in json.loads(get(url(pg)))['responseData']]
        people = sum(int(r['people_total']) for r in rows)
        old = sum(int(v or 0) for r in rows for k, v in r.items() if k.startswith('people_age_') and int(k[11:14]) >= 65)
        return [{'iso3': 'TWN', 'name': 'Taiwan', 'population': people, 'aged_65_plus_pct': round(100 * old / people, 2),
                 'density_per_km2': round(people / 36197, 1), 'period': f'{y}-{m + 1:02d}'}]   # land area 36,197 km² (MOI)
    raise RuntimeError('Taiwan: no month in the last six')


def pacific():
    """Pacific Community (SPC) .Stat, mid-year population estimates (DF_POP_PROJ)."""
    spc = {'CK': ('COK', 'Cook Islands'), 'NU': ('NIU', 'Niue'), 'WF': ('WLF', 'Wallis and Futuna'), 'PN': ('PCN', 'Pitcairn Islands')}
    text = get(f'https://stats-sdmx-disseminate.pacificdata.org/rest/data/SPC,DF_POP_PROJ,3.0/A.{"+".join(spc)}.MIDYEARPOPEST._T._T'
               f'?endPeriod={datetime.date.today().year}&lastNObservations=1&format=csv').decode()
    return [{'iso3': spc[r['GEO_PICT']][0], 'name': spc[r['GEO_PICT']][1], 'population': int(float(r['OBS_VALUE'])),
             'period': r['TIME_PERIOD']} for r in csv.DictReader(io.StringIO(text))]


def aland():
    """Statistics and Research Åland (ÅSUB) PxWeb BE001: 31 December population by single year of age."""
    q = {'query': [{'code': 'år', 'selection': {'filter': 'top', 'values': ['1']}},
                   {'code': 'ålder', 'selection': {'filter': 'all', 'values': ['*']}},
                   {'code': 'kommun', 'selection': {'filter': 'item', 'values': ['MK21']}},
                   {'code': 'kön', 'selection': {'filter': 'item', 'values': ['0']}}], 'response': {'format': 'json-stat2'}}
    d = json.loads(get('https://pxweb.asub.ax/PXWeb/api/v1/sv/Statistik/BE/Befolkningens%20storlek%20och%20struktur/BE001.px',
                       json.dumps(q).encode(), {'Content-Type': 'application/json'}))
    age, v = d['dimension']['ålder']['category']['index'], d['value']
    total = v[age['SSS']]; old = sum(v[i] for k, i in age.items() if k != 'SSS' and int(k) >= 65)
    return [{'iso3': 'ALA', 'name': 'Åland', 'population': total, 'aged_65_plus_pct': round(100 * old / total, 2),
             'period': f"{next(iter(d['dimension']['år']['category']['index']))}-12-31"}]


def jersey():
    """Statistics Jersey, annual population estimates by single year of age and sex (Jersey open data)."""
    meta = json.loads(get('https://opendata.gov.je/api/3/action/package_show?id=jersey-population-estimate-by-age-and-sex-2010-to-present'))
    rows = list(csv.DictReader(io.StringIO(get(meta['result']['resources'][0]['url']).decode('utf-8-sig'))))
    year = max(r['Year'] for r in rows); rows = [r for r in rows if r['Year'] == year]
    n = lambda r: int(r['Male'] or 0) + int(r['Female'] or 0)
    total = sum(n(r) for r in rows); old = sum(n(r) for r in rows if int(r['Age']) >= 65)
    return [{'iso3': 'JEY', 'name': 'Jersey', 'population': total, 'aged_65_plus_pct': round(100 * old / total, 2), 'period': year}]


def norfolk():
    """Australian Bureau of Statistics, estimated resident population (30 June), Norfolk Island = SA2 901041004."""
    rows = list(csv.DictReader(io.StringIO(get('https://data.api.abs.gov.au/rest/data/ABS,ABS_ANNUAL_ERP_ASGS2021,1.2.0/'
                                               'ERP.SA2.901041004.A?format=csv').decode())))
    r = max(rows, key=lambda r: r['TIME_PERIOD'])
    return [{'iso3': 'NFK', 'name': 'Norfolk Island', 'population': int(float(r['OBS_VALUE'])), 'period': f"{r['TIME_PERIOD']}-06-30"}]


def insee():
    """INSEE, legal populations of the overseas collectivities (Saint-Pierre-et-Miquelon, Saint-Barthélemy, Saint-Martin)."""
    # ponytail: the page id is per edition (8680667 = populations of reference 1 Jan 2023, in force 2026); bump it yearly
    t = get('https://www.insee.fr/fr/statistiques/8680667').decode()
    iso = {'975': ('SPM', 'Saint Pierre and Miquelon'), '977': ('BLM', 'Saint-Barthélemy'), '978': ('MAF', 'Saint-Martin')}
    cell = lambda c: re.sub(r'\s+', ' ', html.unescape(re.sub(r'<[^>]+>', '', c))).strip()
    out = []
    for row in re.findall(r'<tr.*?</tr>', t, re.S):
        c = [cell(x) for x in re.findall(r'<t[hd][^>]*>(.*?)</t[hd]>', row, re.S)]
        if len(c) == 3 and c[0] in iso:
            out.append({'iso3': iso[c[0]][0], 'name': iso[c[0]][1], 'population': int(re.sub(r'\D', '', c[2])), 'period': '2023-01-01'})
    if len(out) != 3:
        raise RuntimeError(f'INSEE: expected 3 collectivities, found {len(out)}')
    return out


def guernsey():
    """States of Guernsey Data & Analysis: the population page links the latest 'Supplementary Data (Population up to ...)' XLSX;
    its 'Age' sheet holds quarterly counts by single year of age (the id changes per release, so the link is scraped)."""
    m = re.search(r'href="([^"]*CHttpHandler\.ashx\?id=\d+[^"]*)"[^>]*>(?:\s|<[^>]+>)*Supplementary Data \(Population',
                  get('https://www.gov.gg/population').decode())
    if not m:
        raise RuntimeError('Guernsey: no Supplementary Data link on the population page')
    rows = list(xlsx(urllib.request.urljoin('https://www.gov.gg/', html.unescape(m.group(1))))['Age'].iter_rows(values_only=True))
    total = next(r for r in rows if r[0] == 'TOTAL')
    j = max(i for i, h in enumerate(rows[0]) if str(h).startswith('Count_') and isinstance(total[i], (int, float)))
    old = sum(r[j] for r in rows[1:] if r[0] == '95 and over' or isinstance(r[0], int) and r[0] >= 65)
    d, mth, y = rows[0][j][6:].split('.')
    return [{'iso3': 'GGY', 'name': 'Guernsey', 'population': int(total[j]), 'aged_65_plus_pct': pct(old, total[j]),
             'period': f'{y}-{mth}-{d}'}]


def northern_cyprus():
    """TRNC Statistical Institute (İstatistik Kurumu), population and demography workbook, sheet N3: year-end population by
    five-year age group. These are projections from the 2011 census, not counts."""
    base = 'https://stat.gov.ct.tr'
    try:
        page = get(base + '/TEMEL-%C4%B0STAT%C4%B0ST%C4%B0KLER/N%C3%9CFUS-VE-DEMOGRAF%C4%B0-%C4%B0STAT%C4%B0ST%C4%B0KLER%C4%B0', tries=3).decode()
        links = re.findall(r'href="(/Portals/39/NufusDemo_(\d{4})[^"]*\.xlsx)"', page)
        path = max(links, key=lambda l: l[1])[0]
    except (OSError, ValueError):                                      # page down or relaid out: the file Helen checked
        path = '/Portals/39/NufusDemo_2023_1.xlsx'
    rows = list(xlsx(base + path, tries=3)['N3'].iter_rows(values_only=True))
    years = next(r for r in rows if r[1] == 'Yıllar')
    total = next(r for r in rows if r[0] == 'TOPLAM NÜFUS')
    j = max(i for i, y in enumerate(years) if isinstance(y, int) and isinstance(total[i], (int, float)))
    block = rows[rows.index(total):]
    block = block[:next((k for k, r in enumerate(block[1:], 1) if r[1] == 'Yıllar'), len(block))]   # stop at the male table
    old = sum(r[j] for r in block if r[1] in ('65-69', '70-74', '75+'))
    return [{'iso3': 'CYN', 'name': 'Northern Cyprus', 'population': int(total[j]), 'aged_65_plus_pct': pct(old, total[j]),
             'period': f'{years[j]}-12-31', 'notes': 'Year-end projection from the 2011 census (TRNC Statistical Institute), not a count'}]


def anguilla():
    """Anguilla Statistics Department, '1.1 Population' workbook, Table 1.01.1.0: annual population estimates.
    Ages exist only from the 2011 census, so no 65+ share."""
    rows = list(xlsx('https://statistics.gov.ai/AllDocuments/1.1%20Population.xlsx')['Table 1.01.1.0'].iter_rows(values_only=True))
    years = next(r for r in rows if 1992 in r)
    est = next(r for r in rows if 'Population Estimates' in r)
    j = max(i for i, y in enumerate(years) if re.fullmatch(r'\d{4}', str(y or '').strip()) and isinstance(est[i], (int, float)))
    return [{'iso3': 'AIA', 'name': 'Anguilla', 'population': round(est[j]), 'period': str(years[j]),
             'notes': 'Population estimate (Table 1.01.1.0, Population Estimates); age data in this file is from the 2011 census only'}]


def montserrat():
    """Statistics Department Montserrat, mid-year population estimates by five-year age group (HTML table). The site has carried
    injected spam links, so only this one page is read, nothing on it is followed, and the table must have exactly the
    expected shape: 'Age Group' header, 0-4 ... 85+ rows, a Total row equal to the sum of its age groups."""
    page = get('https://statistics.gov.ms/subjects/social-and-demographic-statistics/population-and-demography/'
               'mid-year-population-estimates-2019-2024/').decode()
    rows = [[text(c) for c in re.findall(r'<t[hd][^>]*>(.*?)</t[hd]>', r, re.S)] for r in re.findall(r'<tr.*?</tr>', page, re.S)]
    years = [c for c in rows[0] if c]
    if not rows or rows[1][0] != 'Age Group' or not all(re.fullmatch(r'\d{4}', y) for y in years):
        raise RuntimeError('Montserrat: table shape changed')
    n = lambda c: int(re.fullmatch(r'[\d,]+', c).group().replace(',', ''))
    ages = {r[0]: n(r[3 * len(years)]) for r in rows[2:] if re.fullmatch(r'\d+-\d+|85\+', r[0])}
    total = n(next(r for r in rows if r[0] == 'Total')[3 * len(years)])
    if len(ages) != 18 or sum(ages.values()) != total:
        raise RuntimeError(f'Montserrat: {len(ages)} age groups summing to {sum(ages.values())}, total {total}')
    old = sum(v for k, v in ages.items() if int(re.match(r'\d+', k).group()) >= 65)
    return [{'iso3': 'MSR', 'name': 'Montserrat', 'population': total, 'aged_65_plus_pct': pct(old, total),
             'period': years[-1], 'notes': 'Mid-year estimate'}]


def saint_helena():
    """St Helena Government Statistics Office 'Population' workbook (monthly): resident on-island population and its 65+ count.
    Plus Tristan da Cunha's islanders from tristandc.com (a sub-place: Natural Earth has no separate Tristan polygon).
    Ascension publishes no figures and is left out."""
    rows = list(xlsx('https://sainthelena.gov.sh/documents/population/')['Month'].iter_rows(values_only=True))
    head = next(r for r in rows if r[0] == 'Code'); res = next(r for r in rows if r[0] == 'POP.RESIDENT')
    old = next(r for r in rows if r[0] == 'POP.RESIDENT.65+')
    j = max(i for i, v in enumerate(res) if i > 5 and isinstance(v, (int, float)) and isinstance(old[i], (int, float)))
    y, m = str(head[j]).split(',')
    # The page's islanders chart: rows [new Date(y, month0, d), on Tristan, overseas]; the latest row is the current count.
    raw = get('https://tristandc.com/population.php').decode('utf-8', 'ignore')
    chart = raw[raw.index('function islandersChart'):raw.index('function populationChart')]
    rows = re.findall(r'\[new Date\((\d{4}), (\d+), (\d+)\), (\d+), (\d+)\]', chart)
    if not rows:
        raise RuntimeError('Tristan da Cunha: islanders chart not found')
    yy, mm, dd, on, away = map(int, max(rows, key=lambda r: tuple(map(int, r[:3]))))
    n = on + away; period = f'{yy}-{mm + 1:02d}-{dd:02d}'
    return [{'iso3': 'SHN', 'name': 'Saint Helena', 'population': int(res[j]), 'aged_65_plus_pct': pct(old[j], res[j]),
             'period': f'{y.strip()}-{int(m):02d}', 'notes': 'Resident on-island population (six-month rule), St Helena island only'},
            {'iso3': 'SHN-TA', 'name': 'Tristan da Cunha', 'population': n, 'on_island': on, 'period': period,
             'notes': 'Tristan da Cunha Islanders (residency status, including those temporarily overseas); a sub-place of '
                      'Saint Helena, Ascension and Tristan da Cunha with no polygon of its own in Natural Earth',
             'publisher': 'Tristan da Cunha Government and Tristan da Cunha Association', 'url': 'https://tristandc.com/population.php'}]


def falklands():
    """Falkland Islands Government Policy Department, 2021 Census tables: Table 1b estimated usual residents; 65+ share from
    Table 2b (usual residents counted, by single year of age)."""
    wb = xlsx('https://www.gov.fk/policy/wp-content/uploads/sites/33/2026/05/2021-Census-Tables.xlsx')
    total = next(r[1] for r in wb['Table 1b'].iter_rows(values_only=True) if r[0] == 'TOTAL POPULATION')
    ages = {r[0]: r[1] for r in wb['Table 2b'].iter_rows(values_only=True) if isinstance(r[0], int) or r[0] == '100+'}
    counted = sum(ages.values()); old = sum(v for k, v in ages.items() if k == '100+' or k >= 65)
    return [{'iso3': 'FLK', 'name': 'Falkland Islands', 'population': int(total), 'aged_65_plus_pct': pct(old, counted),
             'period': '2021', 'notes': '2021 census, estimated usual residents; 65+ share from the usual residents counted'}]


def vatican():
    """Governatorato dello Stato della Città del Vaticano, 'Popolazione' page: inhabitants of the State (citizens and not)."""
    t = text(get('https://www.vaticanstate.va/it/stato-e-governo/note-generali/popolazione.html').decode())
    when = re.search(r'agg\. al (\d+) dicembre (\d{4})', t)
    n = re.search(r'Gli abitanti nello Stato \(cittadini e non\) sono complessivamente (\d+)', t)
    if not (n and when):
        raise RuntimeError('Vatican: inhabitants sentence not found')
    return [{'iso3': 'VAT', 'name': 'Vatican City', 'population': int(n.group(1)), 'period': f'{when.group(2)}-12-{int(when.group(1)):02d}',
             'notes': 'Inhabitants of the State, citizens and non-citizens'}]


def morocco_hcp():
    """Haut-Commissariat au Plan, RGPH 2024 legal population by region: the two regions Morocco administers in Western Sahara.
    Kept as regions, side by side; never summed into a Western Sahara total."""
    rows = list(xlsx('https://hcp.ma/file/242341').worksheets[0].iter_rows(values_only=True))
    out = []
    for name in ('Laâyoune-Sakia El Hamra', 'Dakhla-Oued Ed-Dahab'):
        r = next(r for r in rows if r[0] == f'Région de {name}')
        out.append({'iso3': 'ESH', 'name': name, 'population': int(r[3]), 'period': '2024',
                    'scope': "Moroccan-administered region (Morocco's regional borders, which do not match the territory of Western Sahara)",
                    'status_note': 'Status disputed'})
    return out


def somaliland_manual():
    """Central Statistics Department, Somaliland in Figures (20th edition, November 2022): 'Total Population 4.3 Million
    (Projected From Pess 2014)'. No machine-readable source, so entered by hand."""
    return [{'iso3': 'SOL', 'name': 'Somaliland', 'population': 4_300_000, 'period': '2022',
             'notes': "'Total Population 4.3 Million (Projected From Pess 2014)': a projection from the 2014 Population "
                      'Estimation Survey of Somalia, rounded to 0.1 million',
             'entered_by_hand': '2026-10-10, from the PDF above; re-check yearly'}]


NOLIC = {'licence_note': 'no open licence; facts republished with credit'}

# id: (function, publisher, licence, url people can follow[, extra fields for the file])
SOURCES = {
    'taiwan': (taiwan, 'Ministry of the Interior (Taiwan), Department of Household Registration',
               'https://data.gov.tw/license', 'https://www.ris.gov.tw/app/portal/346'),
    'pacific': (pacific, 'Pacific Community (SPC), Statistics for Development Division',
                'https://pacificdata.org/terms-of-use', 'https://stats.pacificdata.org'),
    'aland': (aland, 'Statistics and Research Åland (ÅSUB)', 'https://www.asub.ax/sv/om-asub/om-asubax/anvandarvillkor', 'https://www.asub.ax/sv/statistik/befolkning'),
    'jersey': (jersey, 'Statistics Jersey, Government of Jersey', 'https://www.gov.je/ServiceManual/opendata/Pages/open-government-licence-jersey-ogl-j-v1-0.aspx',
               'https://opendata.gov.je/dataset/jersey-population-estimate-by-age-and-sex-2010-to-present'),
    'norfolk': (norfolk, 'Australian Bureau of Statistics', 'https://www.abs.gov.au/website-privacy-copyright-and-disclaimer',
                'https://www.abs.gov.au/statistics/people/population/regional-population'),
    'insee': (insee, 'INSEE (Source : Insee)', 'https://www.insee.fr/fr/information/2008466', 'https://www.insee.fr/fr/statistiques/8680667'),
    'guernsey': (guernsey, 'States of Guernsey, Data & Analysis', 'https://www.gov.gg/population', 'https://www.gov.gg/population', NOLIC),
    'northern_cyprus': (northern_cyprus, 'TRNC Statistical Institute (İstatistik Kurumu)', 'https://stat.gov.ct.tr/Terms',
                        'https://stat.gov.ct.tr/TEMEL-%C4%B0STAT%C4%B0ST%C4%B0KLER/N%C3%9CFUS-VE-DEMOGRAF%C4%B0-%C4%B0STAT%C4%B0ST%C4%B0KLER%C4%B0', NOLIC),
    'anguilla': (anguilla, 'Anguilla Statistics Department', 'https://statistics.gov.ai/',
                 'https://statistics.gov.ai/AllDocuments/1.1%20Population.xlsx', NOLIC),
    'montserrat': (montserrat, 'Statistics Department, Ministry of Finance and Economic Management, Montserrat', 'https://statistics.gov.ms/',
                   'https://statistics.gov.ms/subjects/social-and-demographic-statistics/population-and-demography/mid-year-population-estimates-2019-2024/', NOLIC),
    'saint_helena': (saint_helena, 'St Helena Government Statistics Office (Tristan da Cunha: Tristan da Cunha Government and Association)',
                     'https://www.sainthelena.gov.sh/st-helena/statistics/', 'https://sainthelena.gov.sh/documents/population/', NOLIC),
    'falklands': (falklands, 'Falkland Islands Government, Policy Department', 'https://www.gov.fk/policy/',
                  'https://www.gov.fk/policy/wp-content/uploads/sites/33/2026/05/2021-Census-Tables.xlsx', NOLIC),
    'vatican': (vatican, 'Governatorato dello Stato della Città del Vaticano', 'https://www.vaticanstate.va/',
                'https://www.vaticanstate.va/it/stato-e-governo/note-generali/popolazione.html', NOLIC),
    'morocco_hcp': (morocco_hcp, 'Haut-Commissariat au Plan (HCP), Morocco', 'https://www.hcp.ma/Conditions-generales-d-utilisation-Version-1-0_a2194.html',
                    'https://www.hcp.ma/Population-legale-du-Royaume-du-Maroc-repartie-par-regions-provinces-et-prefectures-et-communes-selon-les-resultats-du_a3974.html',
                    {'licence_note': 'CC BY 4.0 for indicators, per HCP terms of use v1.0 (source: HCP, RGPH 2024)'}),
    'somaliland_manual': (somaliland_manual, 'Central Statistics Department, Ministry of Planning and National Development, Somaliland',
                          'https://somalilandcsd.org/', 'https://somalilandcsd.org/wp-content/uploads/2022/12/A5-CSD_SOMALILAND-IN-FIGURES.pdf', NOLIC),
}


def countries():
    """Natural Earth 1:10m country outlines (public domain) as vector tiles, layer `countries`, every field. 1:50m was too coarse
    at city zoom (Macau's shape sat km off the real Macau; Helen, 2026-10-10). Natural Earth draws Taiwan, Hong Kong and Macau
    as their own places, unlike GISCO."""
    src = os.path.join(sources.CACHE, 'ne_10m_admin_0_countries.geojson')
    if not os.path.exists(src):
        os.makedirs(sources.CACHE, exist_ok=True)
        open(src, 'wb').write(get('https://raw.githubusercontent.com/nvkelso/natural-earth-vector/v5.1.2/geojson/ne_10m_admin_0_countries.geojson'))
    out = os.path.join(sources.OUT, 'countries.pmtiles')
    subprocess.run(['tippecanoe', '-o', out, '--force', '-l', 'countries', '-Z0', '-z10', '--detect-shared-borders',
                    '--no-tile-size-limit', '--no-feature-limit', '--quiet', src], check=True)
    return len(json.load(open(src))['features'])


# Sources with no open licence (republished with credit, Helen 2026-10-10) go to a separate dataset, world-places-nc, and only
# the separate world-views-nc site reads them; the open world-places and world-views use open-licence sources only.
NC = {'guernsey', 'northern_cyprus', 'anguilla', 'montserrat', 'saint_helena', 'falklands', 'vatican', 'somaliland_manual'}
NC_OUT = os.path.join(sources.OUT, 'nc')


def write_index(out, ids, failed, extra=None):
    files = sorted(f[:-5] for f in os.listdir(out) if f.endswith('.json') and f[:-5] in ids)
    counts = {f: len(json.load(open(os.path.join(out, f + '.json')))['places']) for f in files}
    idx = {'built': datetime.datetime.now().isoformat(timespec='seconds'), 'counts': {**counts, **(extra or {})},
           'files': [f + '.json' for f in files], 'failed': [x for x in failed if x.split(':')[0] in ids or x.startswith('countries:') and extra]}
    if extra:
        idx['tiles'] = 'countries.pmtiles' if os.path.exists(os.path.join(out, 'countries.pmtiles')) else None
    json.dump(idx, open(os.path.join(out, 'index.json'), 'w'), indent=1)


def keep_published(sid):
    """A source failed (a flaky site): keep its last published file, so one bad fetch does not drop the place. The run still
    lists the failure in index.json. CI starts from an empty folder, so the file comes from the dataset's bucket."""
    out = NC_OUT if sid in NC else sources.OUT
    repo = 'world-places-nc' if sid in NC else 'world-places'
    try:
        open(os.path.join(out, f'{sid}.json'), 'wb').write(get(f'https://huggingface.co/buckets/milleleaf/{repo}/resolve/{sid}.json'))
        print(f'{sid}: kept the published file', file=sys.stderr)
    except Exception as e:
        print(f'{sid}: no published file to keep ({e})', file=sys.stderr)


def main(ids):
    os.makedirs(NC_OUT, exist_ok=True)
    failed = []
    for sid in ids:
        f, publisher, licence, url, *extra = SOURCES[sid]
        try:
            places = f()
        except Exception as e:
            failed.append(f'{sid}: {e}'); print(f'{sid}: FAILED {e}', file=sys.stderr)
            keep_published(sid); continue
        out = NC_OUT if sid in NC else sources.OUT
        json.dump({'source': sid, 'publisher': publisher, 'url': url, 'licence': licence, **(extra[0] if extra else {}), 'places': places},
                  open(os.path.join(out, f'{sid}.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
        print(sid, [(p['iso3'], p['population'], p['period']) for p in places])
    for stale in [f for f in os.listdir(sources.OUT) if f[:-5] in NC]:      # an older bake wrote nc sources to the open folder
        os.remove(os.path.join(sources.OUT, stale))
    extra = {}
    try:
        extra['countries'] = countries()
    except Exception as e:
        failed.append(f'countries: {e}'); print(f'countries: FAILED {e}', file=sys.stderr)
    write_index(sources.OUT, set(SOURCES) - NC, failed, extra)
    write_index(NC_OUT, NC, failed)
    if len([f for f in failed if not f.startswith('countries:')]) == len(ids):
        sys.exit('every source failed')


if __name__ == '__main__':
    main(sys.argv[1:] or list(SOURCES))
